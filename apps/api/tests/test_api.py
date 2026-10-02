import io
import time
from zipfile import ZipFile

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from convertbox_api.core import JobStatus, can_transition
from convertbox_api.main import app


def wait_for_job(client: TestClient, job_id: str) -> dict:
    for _ in range(120):
        state = client.get(f"/jobs/{job_id}").json()
        if state["status"] in ("COMPLETED", "FAILED"):
            return state
        time.sleep(0.1)
    raise AssertionError("Job did not finish")


def submit(client: TestClient, data: bytes, name: str, output: str, **fields: str) -> tuple[dict, bytes]:
    response = client.post("/jobs", files=[("files", (name, data))], data={"output": output, **fields})
    assert response.status_code == 202, response.text
    job_id = response.json()["id"]
    state = wait_for_job(client, job_id)
    assert state["status"] == "COMPLETED", state
    result = client.get(f"/jobs/{job_id}/download")
    assert result.status_code == 200
    return state, result.content


def test_capabilities_and_transitions() -> None:
    client = TestClient(app)
    assert client.get("/health").json() == {"status": "ok"}
    capabilities = client.get("/capabilities").json()
    assert capabilities["version"] == 3
    assert any(item["input"] == "pdf" and "png" in item["outputs"] for item in capabilities["server"])
    assert can_transition(JobStatus.QUEUED, JobStatus.PROCESSING)
    assert not can_transition(JobStatus.COMPLETED, JobStatus.PROCESSING)


def test_image_pdf_office_and_download() -> None:
    image = Image.new("RGB", (12, 8), "red")
    image_bytes = io.BytesIO()
    image.save(image_bytes, "PNG")
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "ConvertBox text PDF")
    pdf_bytes = pdf.tobytes()
    pdf.close()
    word = Document()
    word.add_paragraph("ConvertBox Word document")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    with TestClient(app) as client:
        _, converted = submit(client, image_bytes.getvalue(), "sample.png", "pdf")
        with pymupdf.open(stream=converted, filetype="pdf") as result:
            assert len(result) == 1
        _, images_zip = submit(client, pdf_bytes, "sample.pdf", "png", settings='{"dpi":72}')
        with ZipFile(io.BytesIO(images_zip)) as archive:
            rendered = Image.open(io.BytesIO(archive.read("page_1.png")))
            assert rendered.size == (595, 842)
        _, text = submit(client, pdf_bytes, "sample.pdf", "txt")
        assert b"ConvertBox text PDF" in text
        _, docx = submit(client, pdf_bytes, "sample.pdf", "docx")
        recovered = Document(io.BytesIO(docx))
        assert "ConvertBox text PDF" in recovered.paragraphs[0].text
        _, office_pdf = submit(client, word_bytes.getvalue(), "sample.docx", "pdf")
        with pymupdf.open(stream=office_pdf, filetype="pdf") as result:
            assert "ConvertBox Word document" in result[0].get_text()
        response = client.post("/jobs", files=[("files", ("fake.png", b"not a PNG"))], data={"output": "pdf"})
        assert response.status_code == 400


def test_pdf_merge_split_and_delete() -> None:
    first = pymupdf.open()
    first.new_page().insert_text((72, 72), "First")
    second = pymupdf.open()
    second.new_page().insert_text((72, 72), "Second")
    with TestClient(app) as client:
        response = client.post("/jobs", files=[
            ("files", ("one.pdf", first.tobytes())),
            ("files", ("two.pdf", second.tobytes())),
        ], data={"output": "pdf", "operation": "merge"})
        assert response.status_code == 202, response.text
        state = wait_for_job(client, response.json()["id"])
        assert state["status"] == "COMPLETED"
        merged = client.get(f"/jobs/{state['id']}/download").content
        with pymupdf.open(stream=merged, filetype="pdf") as result:
            assert len(result) == 2
        _, split = submit(client, merged, "merged.pdf", "pdf", operation="split")
        with ZipFile(io.BytesIO(split)) as archive:
            assert archive.namelist() == ["page_1.pdf", "page_2.pdf"]
        assert client.delete(f"/jobs/{state['id']}").status_code == 204
        assert client.get(f"/jobs/{state['id']}").status_code == 404
    first.close()
    second.close()
