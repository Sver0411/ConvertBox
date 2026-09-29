import io
import json
import subprocess
from pathlib import Path
from zipfile import ZipFile

import pymupdf
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from convertbox_api.converters import ImageConverter, MediaConverter, OfficeConverter, PdfConverter
from convertbox_api.core import ConversionRequest
from convertbox_api.detection import detect_file

FIXTURES = Path(__file__).resolve().parents[3] / "tests" / "fixtures"


def request(source: Path, target: Path, source_format: str, target_format: str, **settings: int | str) -> ConversionRequest:
    return ConversionRequest(source, target, source_format, target_format, settings)


def test_heic_avif_bmp_gif_and_resize(tmp_path: Path) -> None:
    converter = ImageConverter()
    for source_format in ("heic", "avif"):
        source = FIXTURES / f"sample.{source_format}"
        assert detect_file(source, source.name) == source_format
        if source_format == "heic":
            assert detect_file(source, "photo.heif") == "heic"
        output = tmp_path / f"{source_format}.jpg"
        converter.convert(request(source, output, source_format, "jpg", width=8))
        with Image.open(output) as image:
            assert image.format == "JPEG"
            assert image.size == (8, 6)
    for source_format in ("bmp", "gif"):
        source = tmp_path / f"sample.{source_format}"
        Image.new("RGB", (10, 5), "red").save(source, source_format.upper())
        assert detect_file(source, source.name) == source_format
        output = tmp_path / f"{source_format}.webp"
        converter.convert(request(source, output, source_format, "webp"))
        with Image.open(output) as image:
            assert image.format == "WEBP"


def test_exif_keep_and_remove(tmp_path: Path) -> None:
    source = tmp_path / "metadata.jpg"
    image = Image.new("RGB", (16, 12), "green")
    exif = Image.Exif()
    exif[270] = "ConvertBox camera note"
    image.save(source, "JPEG", exif=exif)
    converter = ImageConverter()
    for keep in (False, True):
        output = tmp_path / f"metadata_{keep}.jpg"
        converter.convert(ConversionRequest(source, output, "jpg", "jpg", {"keep_metadata": keep}))
        with Image.open(output) as result:
            assert result.getexif().get(270) == ("ConvertBox camera note" if keep else None)


def test_alpha_orientation_and_invalid_image(tmp_path: Path) -> None:
    source = tmp_path / "透明🌿.png"
    Image.new("RGBA", (4, 2), (255, 0, 0, 0)).save(source)
    assert detect_file(source, source.name) == "png"
    output = tmp_path / "transparent.png"
    ImageConverter().convert(request(source, output, "png", "png"))
    with Image.open(output) as result:
        assert result.size == (4, 2)
        assert result.getpixel((0, 0))[3] == 0

    rotated = tmp_path / "回転.jpg"
    exif = Image.Exif()
    exif[274] = 6
    Image.new("RGB", (4, 2), "blue").save(rotated, "JPEG", exif=exif)
    oriented = tmp_path / "oriented.png"
    ImageConverter().convert(request(rotated, oriented, "jpg", "png"))
    with Image.open(oriented) as result:
        assert result.size == (2, 4)

    corrupt = tmp_path / "broken.jpg"
    corrupt.write_bytes(b"\xff\xd8\xff\xe0" + b"broken")
    assert detect_file(corrupt, corrupt.name) == "jpg"
    try:
        ImageConverter().convert(request(corrupt, tmp_path / "bad.png", "jpg", "png"))
    except Exception:
        pass
    else:
        raise AssertionError("Corrupt JPEG was converted")


def test_pdf_pages_rotate_compress(tmp_path: Path) -> None:
    source = tmp_path / "pages.pdf"
    with pymupdf.open() as document:
        for text in ("one", "two", "three"):
            document.new_page().insert_text((72, 72), text)
        document.save(source)
    converter = PdfConverter()
    selected = tmp_path / "selected.zip"
    converter.convert(ConversionRequest(source, selected, "pdf", "png", {"pages": "1,3", "dpi": 72}))
    with ZipFile(selected) as archive:
        assert archive.namelist() == ["page_1.png", "page_3.png"]
        with Image.open(io.BytesIO(archive.read("page_1.png"))) as image:
            assert image.size == (595, 842)
    rotated = tmp_path / "rotated.pdf"
    converter.convert(ConversionRequest(source, rotated, "pdf", "pdf", {"pages": "2", "rotation": 90}, operation="rotate"))
    with pymupdf.open(rotated) as result:
        assert [page.rotation for page in result] == [0, 90, 0]
    compressed = tmp_path / "compressed.pdf"
    converter.convert(ConversionRequest(source, compressed, "pdf", "pdf", {}, operation="compress"))
    with pymupdf.open(compressed) as result:
        assert len(result) == 3
    split = tmp_path / "groups.zip"
    converter.convert(ConversionRequest(source, split, "pdf", "pdf", {"pages": "1-2;3"}, operation="split"))
    with ZipFile(split) as archive:
        assert archive.namelist() == ["pages_1-2.pdf", "pages_3.pdf"]
        with pymupdf.open(stream=archive.read("pages_1-2.pdf"), filetype="pdf") as result:
            assert len(result) == 2


def test_images_to_pdf_layout_and_order(tmp_path: Path) -> None:
    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    Image.new("RGB", (80, 40), "red").save(first)
    Image.new("RGB", (40, 80), "blue").save(second)
    output = tmp_path / "combined.pdf"
    ImageConverter().convert(ConversionRequest(first, output, "png", "pdf", {"page_size": "a4", "orientation": "landscape", "margin": "small"}, input_paths=(first, second)))
    with pymupdf.open(output) as result:
        assert len(result) == 2
        assert all(page.rect.width > page.rect.height for page in result)
        assert result[0].get_images() and result[1].get_images()


def test_spreadsheet_and_presentation_to_pdf(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "ConvertBox spreadsheet"
    spreadsheet = tmp_path / "sample.xlsx"
    workbook.save(spreadsheet)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(100000, 100000, 3000000, 500000)
    box.text = "ConvertBox slides"
    slides = tmp_path / "sample.pptx"
    presentation.save(slides)
    converter = OfficeConverter()
    for source, kind in ((spreadsheet, "xlsx"), (slides, "pptx"), *((FIXTURES / f"sample.{kind}", kind) for kind in ("doc", "odt", "xls", "ods", "ppt", "odp"))):
        assert detect_file(source, source.name) == kind
        output = tmp_path / f"{kind}.pdf"
        converter.convert(request(source, output, kind, "pdf"))
        with pymupdf.open(output) as result:
            assert len(result) >= 1


def test_audio_and_video_variants(tmp_path: Path) -> None:
    converter = MediaConverter()
    for source, source_format, target_format in (
        (FIXTURES / "sample.wav", "wav", "flac"),
        (FIXTURES / "sample.wav", "wav", "ogg"),
        (FIXTURES / "sample.mp4", "mp4", "webm"),
        (FIXTURES / "sample.mp4", "mp4", "mkv"),
    ):
        assert detect_file(source, source.name) == source_format
        output = tmp_path / f"result_{source_format}.{target_format}"
        converter.convert(request(source, output, source_format, target_format))
        result = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1", str(output)], capture_output=True, text=True, check=True)
        assert "duration=" in result.stdout
        streams = subprocess.run(["ffprobe", "-v", "error", "-show_streams", "-of", "json", str(output)], capture_output=True, text=True, check=True)
        kinds = {stream["codec_type"] for stream in json.loads(streams.stdout)["streams"]}
        assert ("video" if source_format == "mp4" else "audio") in kinds
